import io
import os
import uuid
import qrcode
from datetime import datetime, date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import get_settings
from app.models.competition import Competition
from app.models.result import Result, ResultStatus
from app.models.judge import JudgeAssignment
from app.models.rating import TeamRoster
from app.models.document import Certificate
from app.services.evsk_service import evaluate_evsk_rank

settings = get_settings()

CYRILLIC_FONT = "Helvetica"
# Кандидаты (обычный, жирный): Windows, Linux/Docker (Liberation = метрически совместим с Arial), DejaVu
_FONT_CANDIDATES = [
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
]
try:
    for regular, bold in _FONT_CANDIDATES:
        if os.path.exists(regular):
            pdfmetrics.registerFont(TTFont("Arial", regular))
            pdfmetrics.registerFont(TTFont("Arial-Bold", bold if os.path.exists(bold) else regular))
            CYRILLIC_FONT = "Arial"
            break
except Exception:
    pass


class DocumentService:
    @staticmethod
    def generate_protocol_docx(comp_data: dict, results_data: list[dict], judges_data: list[dict]) -> io.BytesIO:
        doc = Document()

        p_header = doc.add_paragraph()
        p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p_header.add_run("ФЕДЕРАЦИЯ СПОРТИВНОГО ПРОГРАММИРОВАНИЯ ОРЕНБУРГСКОЙ ОБЛАСТИ\n")
        r1.bold = True
        r1.font.size = Pt(13)

        r2 = p_header.add_run(f"ОФИЦИАЛЬНЫЙ ИТОГОВЫЙ ПРОТОКОЛ\n«{comp_data.get('title', '')}»\n")
        r2.bold = True
        r2.font.size = Pt(14)

        p_meta = doc.add_paragraph()
        p_meta.add_run("Дисциплина: ").bold = True
        p_meta.add_run(f"{comp_data.get('discipline_name', '—')}\n")
        p_meta.add_run("Сроки проведения: ").bold = True
        p_meta.add_run(f"{comp_data.get('start_date', '—')} — {comp_data.get('end_date', '—')}\n")
        p_meta.add_run("Место проведения: ").bold = True
        p_meta.add_run(f"{comp_data.get('location', 'г. Оренбург')}\n")
        p_meta.add_run("Уровень: ").bold = True
        p_meta.add_run(f"{comp_data.get('level', '—')}\n")

        doc.add_paragraph()

        # Results table
        table = doc.add_table(rows=1, cols=7)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'

        headers = ["Место", "ФИО спортсмена", "Организация", "Город", "Разряд", "Баллы", "Статус"]
        for i, title in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = title
            cell.paragraphs[0].runs[0].bold = True
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        for r in results_data:
            row = table.add_row().cells
            row[0].text = str(r.get("place") or "—")
            row[1].text = r.get("athlete_name", "—")
            row[2].text = r.get("athlete_organization", "—")
            row[3].text = r.get("athlete_municipality", "—")
            row[4].text = r.get("rank", "б/р")
            row[5].text = str(r.get("score") if r.get("score") is not None else "—")
            row[6].text = r.get("status", "normal")

        doc.add_paragraph()

        # Judges
        p_judges = doc.add_paragraph()
        p_judges.add_run("ГЛАВНАЯ СУДЕЙСКАЯ КОЛЛЕГИЯ:\n").bold = True
        if judges_data:
            for j in judges_data:
                p_judges.add_run(f"• {j.get('position', 'Судья')}: {j.get('judge_name', '—')} ({j.get('category', 'б/к')})\n")
        else:
            p_judges.add_run("• Главный судья: __________________ (подпись)\n")
            p_judges.add_run("• Главный секретарь: ______________ (подпись)\n")

        p_footer = doc.add_paragraph()
        p_footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_footer.add_run(f"\nДата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}")

        stream = io.BytesIO()
        doc.save(stream)
        stream.seek(0)
        return stream

    @staticmethod
    def generate_protocol_pdf(comp_data: dict, results_data: list[dict], judges_data: list[dict]) -> io.BytesIO:
        stream = io.BytesIO()
        doc = SimpleDocTemplate(stream, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        styles = getSampleStyleSheet()

        font_name = CYRILLIC_FONT
        font_bold = "Arial-Bold" if font_name == "Arial" else font_name

        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Normal'],
            fontName=font_bold,
            fontSize=12,
            alignment=1,
            spaceAfter=10
        )
        normal_style = ParagraphStyle(
            'NormalStyle',
            parent=styles['Normal'],
            fontName=font_name,
            fontSize=9,
            spaceAfter=4
        )
        cell_style = ParagraphStyle(
            'CellStyle',
            parent=styles['Normal'],
            fontName=font_name,
            fontSize=8,
            leading=10
        )
        cell_bold = ParagraphStyle(
            'CellBold',
            parent=styles['Normal'],
            fontName=font_bold,
            fontSize=8,
            leading=10
        )

        story = []
        story.append(Paragraph("ФЕДЕРАЦИЯ СПОРТИВНОГО ПРОГРАММИРОВАНИЯ ОРЕНБУРГСКОЙ ОБЛАСТИ", title_style))
        story.append(Paragraph(f"ОФИЦИАЛЬНЫЙ ИТОГОВЫЙ ПРОТОКОЛ<br/>«{comp_data.get('title', '')}»", title_style))
        story.append(Spacer(1, 10))

        meta_text = (
            f"<b>Дисциплина:</b> {comp_data.get('discipline_name', '—')} &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"<b>Сроки:</b> {comp_data.get('start_date', '—')} — {comp_data.get('end_date', '—')}<br/>"
            f"<b>Место проведения:</b> {comp_data.get('location', 'г. Оренбург')} &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"<b>Уровень:</b> {comp_data.get('level', '—')}"
        )
        story.append(Paragraph(meta_text, normal_style))
        story.append(Spacer(1, 12))

        data = [
            [
                Paragraph("<b>Место</b>", cell_bold),
                Paragraph("<b>ФИО спортсмена</b>", cell_bold),
                Paragraph("<b>Организация</b>", cell_bold),
                Paragraph("<b>Разряд</b>", cell_bold),
                Paragraph("<b>Баллы</b>", cell_bold),
                Paragraph("<b>Статус</b>", cell_bold)
            ]
        ]

        for r in results_data:
            data.append([
                Paragraph(str(r.get("place") or "—"), cell_style),
                Paragraph(r.get("athlete_name", "—"), cell_style),
                Paragraph(r.get("athlete_organization", "—"), cell_style),
                Paragraph(r.get("rank", "б/р"), cell_style),
                Paragraph(str(r.get("score") if r.get("score") is not None else "—"), cell_style),
                Paragraph(r.get("status", "normal"), cell_style)
            ])

        table = Table(data, colWidths=[1.5*cm, 5.5*cm, 4.5*cm, 2.0*cm, 2.0*cm, 2.0*cm])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0d47a1')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(table)
        story.append(Spacer(1, 15))
        story.append(Paragraph(f"Протокол сформирован: {datetime.now().strftime('%d.%m.%Y %H:%M')}", normal_style))

        doc.build(story)
        stream.seek(0)
        return stream

    @staticmethod
    def generate_roster_docx(roster_data: dict, main_members: list[dict], reserve_members: list[dict]) -> io.BytesIO:
        doc = Document()

        p_header = doc.add_paragraph()
        p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_hdr = p_header.add_run(
            "МИНИСТЕРСТВО ФИЗИЧЕСКОЙ КУЛЬТУРЫ И СПОРТА ОРЕНБУРГСКОЙ ОБЛАСТИ\n"
            "РОО «ФЕДЕРАЦИЯ СПОРТИВНОГО ПРОГРАММИРОВАНИЯ ОРЕНБУРГСКОЙ ОБЛАСТИ»\n\n"
            "ПРОЕКТ ПРИКАЗА\n\n"
        )
        r_hdr.bold = True
        r_hdr.font.size = Pt(13)

        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_t = p_title.add_run(
            f"«____» ____________ {datetime.now().year} г.                                                  № ______-п\n\n"
            f"Об утверждении списка кандидатов в спортивную сборную команду\n"
            f"Оренбургской области по спортивному программированию на сезон {roster_data.get('season', '')}\n\n"
        )
        r_t.bold = True

        p_preamble = doc.add_paragraph()
        p_preamble.paragraph_format.first_line_indent = Inches(0.5)
        p_preamble.add_run(
            "В соответствии с Федеральным законом от 04.12.2007 № 329-ФЗ «О физической культуре и спорте в "
            "Российской Федерации», правилами вида спорта «спортивное программирование», утвержденными приказом "
            "Минспорта России от 11.08.2023 № 584, а также на основании официального рейтинга спортсменов:\n\n"
            "ПРИКАЗЫВАЮ:\n"
            f"1. Утвердить прилагаемый список кандидатов в спортивную сборную команду Оренбургской области "
            f"по виду спорта «спортивное программирование» (дисциплина: «{roster_data.get('discipline_name', 'Спортивное программирование')}»).\n"
            "2. Контроль за исполнением настоящего приказа оставляю за собой.\n\n"
        )

        doc.add_paragraph().add_run("ОСНОВНОЙ СОСТАВ:").bold = True
        table_main = doc.add_table(rows=1, cols=6)
        table_main.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_main.style = 'Table Grid'
        headers = ["№", "ФИО спортсмена", "Дата рожд.", "Разряд", "Организация / Город", "Рейтинг"]
        for i, h in enumerate(headers):
            cell = table_main.rows[0].cells[i]
            cell.text = h
            cell.paragraphs[0].runs[0].bold = True

        for idx, m in enumerate(main_members, start=1):
            row = table_main.add_row().cells
            row[0].text = str(idx)
            row[1].text = m.get("athlete_name", "—")
            row[2].text = m.get("birth_date", "—")
            row[3].text = m.get("rank", "б/р")
            row[4].text = f"{m.get('organization_name', '—')} ({m.get('municipality', 'Оренбург')})"
            row[5].text = str(m.get("rating_score", 0))

        doc.add_paragraph()
        doc.add_paragraph().add_run("РЕЗЕРВНЫЙ СОСТАВ:").bold = True
        table_res = doc.add_table(rows=1, cols=6)
        table_res.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_res.style = 'Table Grid'
        for i, h in enumerate(headers):
            cell = table_res.rows[0].cells[i]
            cell.text = h
            cell.paragraphs[0].runs[0].bold = True

        for idx, m in enumerate(reserve_members, start=1):
            row = table_res.add_row().cells
            row[0].text = str(idx)
            row[1].text = m.get("athlete_name", "—")
            row[2].text = m.get("birth_date", "—")
            row[3].text = m.get("rank", "б/р")
            row[4].text = f"{m.get('organization_name', '—')} ({m.get('municipality', 'Оренбург')})"
            row[5].text = str(m.get("rating_score", 0))

        doc.add_paragraph()
        p_sign = doc.add_paragraph()
        p_sign.add_run("\nПрезидент РОО «ФСП Оренбургской области»      ______________ / А. В. Фадеев /")

        stream = io.BytesIO()
        doc.save(stream)
        stream.seek(0)
        return stream

    @staticmethod
    def generate_statistics_xlsx(stats_data: dict) -> io.BytesIO:
        wb = Workbook()

        header_fill = PatternFill(start_color="1A237E", end_color="1A237E", fill_type="solid")
        accent_fill = PatternFill(start_color="E8EAF6", end_color="E8EAF6", fill_type="solid")
        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        font_bold = Font(name="Calibri", size=10, bold=True)
        font_regular = Font(name="Calibri", size=10)
        border_thin = Border(
            left=Side(style='thin', color='B0BEC5'),
            right=Side(style='thin', color='B0BEC5'),
            top=Side(style='thin', color='B0BEC5'),
            bottom=Side(style='thin', color='B0BEC5')
        )

        # Sheet 1: Disciplines Dynamics
        ws1 = wb.active
        ws1.title = "Динамика по дисциплинам"

        ws1.merge_cells("A1:G1")
        ws1["A1"] = "ЧИСЛЕННОСТЬ ЗАНИМАЮЩИХСЯ СПОРТИВНЫМ ПРОГРАММИРОВАНИЕМ В ОРЕНБУРГСКОЙ ОБЛАСТИ (2024–2026 гг.)"
        ws1["A1"].font = Font(name="Calibri", size=14, bold=True, color="1A237E")
        ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws1.row_dimensions[1].height = 30

        headers_ws1 = ["№", "Дисциплина вида спорта", "2024 год", "2025 год", "2026 год", "Прирост (чел)", "Динамика (%)"]
        for col_num, h_text in enumerate(headers_ws1, 1):
            cell = ws1.cell(row=3, column=col_num, value=h_text)
            cell.font = font_header
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border_thin

        disciplines = stats_data.get("disciplines", [])
        for idx, d in enumerate(disciplines, start=1):
            row_idx = idx + 3
            ws1.cell(row=row_idx, column=1, value=idx).alignment = Alignment(horizontal="center")
            ws1.cell(row=row_idx, column=2, value=d.get("name", d.get("code", "")))
            ws1.cell(row=row_idx, column=3, value=d.get("y24", 0)).number_format = '#,##0'
            ws1.cell(row=row_idx, column=4, value=d.get("y25", 0)).number_format = '#,##0'
            ws1.cell(row=row_idx, column=5, value=d.get("y26", 0)).number_format = '#,##0'
            ws1.cell(row=row_idx, column=6, value=f"=E{row_idx}-C{row_idx}").number_format = '#,##0'
            ws1.cell(row=row_idx, column=7, value=f"=(E{row_idx}-C{row_idx})/C{row_idx}").number_format = '0.0%'

            for c in range(1, 8):
                ws1.cell(row=row_idx, column=c).font = font_regular
                ws1.cell(row=row_idx, column=c).border = border_thin

        tot_row = len(disciplines) + 4
        ws1.cell(row=tot_row, column=1, value="")
        ws1.cell(row=tot_row, column=2, value="ИТОГО ПО ВСЕМ ДИСЦИПЛИНАМ").font = font_bold
        ws1.cell(row=tot_row, column=3, value=f"=SUM(C4:C{tot_row-1})").font = font_bold
        ws1.cell(row=tot_row, column=4, value=f"=SUM(D4:D{tot_row-1})").font = font_bold
        ws1.cell(row=tot_row, column=5, value=f"=SUM(E4:E{tot_row-1})").font = font_bold
        ws1.cell(row=tot_row, column=6, value=f"=E{tot_row}-C{tot_row}").font = font_bold
        ws1.cell(row=tot_row, column=7, value=f"=(E{tot_row}-C{tot_row})/C{tot_row}").font = font_bold
        ws1.cell(row=tot_row, column=7).number_format = '0.0%'

        for c in range(1, 8):
            ws1.cell(row=tot_row, column=c).fill = accent_fill
            ws1.cell(row=tot_row, column=c).border = border_thin

        ws1.column_dimensions['A'].width = 6
        ws1.column_dimensions['B'].width = 44
        ws1.column_dimensions['C'].width = 14
        ws1.column_dimensions['D'].width = 14
        ws1.column_dimensions['E'].width = 14
        ws1.column_dimensions['F'].width = 16
        ws1.column_dimensions['G'].width = 16

        # Sheet 2: Organizations Breakdown
        ws2 = wb.create_sheet(title="По организациям")
        ws2.append(["Организация", "Тип", "Муниципалитет", "2024 год", "2025 год", "2026 год"])
        for c in range(1, 7):
            ws2.cell(row=1, column=c).font = font_header
            ws2.cell(row=1, column=c).fill = header_fill

        orgs = stats_data.get("organizations", [])
        for o in orgs:
            ws2.append([o.get("name", ""), o.get("type", ""), o.get("mun", ""), o.get("y24", 0), o.get("y25", 0), o.get("y26", 0)])

        for row in ws2.iter_rows(min_row=2, max_row=len(orgs)+1, min_col=1, max_col=6):
            for cell in row:
                cell.border = border_thin
                cell.font = font_regular

        ws2.column_dimensions['A'].width = 50
        ws2.column_dimensions['B'].width = 16
        ws2.column_dimensions['C'].width = 18
        ws2.column_dimensions['D'].width = 14
        ws2.column_dimensions['E'].width = 14
        ws2.column_dimensions['F'].width = 14

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    @staticmethod
    def generate_certificate_pdf(
        cert_number: str,
        athlete_name: str,
        comp_title: str,
        discipline_name: str,
        comp_date: str,
        place: int | None,
        verify_url: str
    ) -> io.BytesIO:
        # Generate QR
        qr = qrcode.QRCode(version=1, error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=10, border=2)
        qr.add_data(verify_url)
        qr.make(fit=True)
        img_qr = qr.make_image(fill_color="black", back_color="white")

        qr_stream = io.BytesIO()
        img_qr.save(qr_stream, format='PNG')
        qr_stream.seek(0)

        stream = io.BytesIO()
        doc = SimpleDocTemplate(stream, pagesize=A4, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        styles = getSampleStyleSheet()

        font_name = CYRILLIC_FONT
        font_bold = "Arial-Bold" if font_name == "Arial" else font_name

        title_style = ParagraphStyle(
            'CertTitle',
            parent=styles['Normal'],
            fontName=font_bold,
            fontSize=18,
            leading=22,
            alignment=1,
            textColor=colors.HexColor('#0d47a1')
        )
        sub_style = ParagraphStyle(
            'CertSub',
            parent=styles['Normal'],
            fontName=font_bold,
            fontSize=12,
            leading=16,
            alignment=1,
            spaceAfter=15,
            textColor=colors.HexColor('#37474f')
        )
        body_style = ParagraphStyle(
            'CertBody',
            parent=styles['Normal'],
            fontName=font_name,
            fontSize=11,
            leading=16,
            spaceAfter=10
        )
        center_bold = ParagraphStyle(
            'CenterBold',
            parent=styles['Normal'],
            fontName=font_bold,
            fontSize=14,
            leading=18,
            alignment=1,
            spaceAfter=12
        )

        story = []
        story.append(Paragraph("РОО «ФЕДЕРАЦИЯ СПОРТИВНОГО ПРОГРАММИРОВАНИЯ<br/>ОРЕНБУРГСКОЙ ОБЛАСТИ»", sub_style))
        story.append(Spacer(1, 10))
        story.append(Paragraph("СПРАВКА ОБ УЧАСТИИ В СОРЕВНОВАНИЯХ", title_style))
        story.append(Paragraph(f"№ {cert_number}", sub_style))
        story.append(Spacer(1, 15))

        story.append(Paragraph("Настоящая справка выдана:", body_style))
        story.append(Paragraph(f"<b>{athlete_name}</b>", center_bold))

        place_text = f"Занятое место: <b>{place}</b>" if place else "Статус участия: подтверждено"

        text = (
            f"в том, что он(а) действительно принял(а) участие в официальном соревновании "
            f"по виду спорта «спортивное программирование»:<br/><br/>"
            f"<b>«{comp_title}»</b><br/>"
            f"Дисциплина: <b>{discipline_name}</b><br/>"
            f"Сроки проведения: <b>{comp_date}</b><br/><br/>"
            f"Результат выступления: {place_text}<br/><br/>"
            f"Справка выдана для подтверждения участия, оформления спортивного портфолио и мер поощрения."
        )
        story.append(Paragraph(text, body_style))
        story.append(Spacer(1, 25))

        qr_img = RLImage(qr_stream, width=3.2*cm, height=3.2*cm)
        sign_text = (
            "<b>Президент Федерации</b><br/>"
            "РОО «ФСП Оренбургской области»<br/><br/>"
            "_________________ / Фадеев А. В. /<br/><br/>"
            f"<i>Дата выдачи: {datetime.now().strftime('%d.%m.%Y')}</i><br/>"
            f"<font size=8 color=gray>Для проверки подлинности отсканируйте QR-код</font>"
        )
        sign_para = Paragraph(sign_text, body_style)

        bottom_table = Table([[qr_img, sign_para]], colWidths=[4.0*cm, 12.0*cm])
        bottom_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ALIGN', (0, 0), (0, 0), 'CENTER'),
        ]))
        story.append(bottom_table)

        doc.build(story)
        stream.seek(0)
        return stream
